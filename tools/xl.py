# -*- coding: utf-8 -*-
"""Ortak .xlsx biçimlendirme yardımcıları — referans hakediş dosyasının düzeni."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

NAVY   = "FF16273D"; NAVY2 = "FF24405F"; COPPER = "FFB87333"
LIGHT  = "FFF4F5F7"; CREAM = "FFFDF3EC"; WHITE  = "FFFFFFFF"
RED    = "FFC8322B"; GREEN = "FF2E7D5B"; GREY   = "FF8A8F98"
INK    = "FF1C1C1C"; BOS   = "FFFFF6E8"      # doldurulacak hücre
GRUPBG = "FFE7EAEE"

_thin = Side(style="thin", color="FFD5D8DC")
_med  = Side(style="medium", color="FF16273D")
BOX   = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
BOX_T = Border(left=_thin, right=_thin, top=_med,  bottom=_thin)

TL   = '#,##0.00 "TL"'
TL0  = '#,##0 "TL"'
NUM  = '#,##0.00'
NUM3 = '#,##0.000'
PCT  = '0%'

def st(c, b=False, sz=10, col=INK, fill=None, al="left", wrap=False, fmt=None,
       border=True, italic=False):
    c.font = Font(name="Calibri", bold=b, size=sz, color=col, italic=italic)
    if fill: c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal=al, vertical="center", wrap_text=wrap)
    if border: c.border = BOX
    if fmt: c.number_format = fmt
    return c

def baslik(ws, satir, metin, alt=None, genislik=12):
    ws.merge_cells(start_row=satir, start_column=1, end_row=satir, end_column=genislik)
    st(ws.cell(satir, 1, metin), True, 13, WHITE, NAVY, "left")
    ws.row_dimensions[satir].height = 24
    if alt:
        ws.merge_cells(start_row=satir+1, start_column=1, end_row=satir+1, end_column=genislik)
        st(ws.cell(satir+1, 1, alt), False, 9, "FF6B7078", LIGHT, "left")
        ws.row_dimensions[satir+1].height = 15
        return satir+2
    return satir+1

def tablo_basligi(ws, satir, basliklar, yukseklik=30):
    for j, x in enumerate(basliklar, 1):
        st(ws.cell(satir, j, x), True, 9, WHITE, NAVY2, "center", True)
    ws.row_dimensions[satir].height = yukseklik
    return satir+1

def genislikler(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x

def grup_satiri(ws, satir, kod, ad, n_sut, fill=GRUPBG):
    st(ws.cell(satir, 1, kod), True, 10, NAVY, fill, "center")
    ws.merge_cells(start_row=satir, start_column=2, end_row=satir, end_column=n_sut)
    st(ws.cell(satir, 2, ad), True, 10, NAVY, fill, "left")
    for j in range(2, n_sut+1): ws.cell(satir, j).border = BOX
    ws.row_dimensions[satir].height = 18
    return satir+1

def dondur(ws, hucre): ws.freeze_panes = hucre

def imza_blogu(ws, satir, sutun=1, genis=4):
    for i, (bas, kim) in enumerate((("HAZIRLAYAN", "Proje müellifi"),
                                    ("KONTROL", "Yapı denetim / mimar"),
                                    ("ONAY", "İşveren"))):
        c0 = sutun + i*genis
        ws.merge_cells(start_row=satir, start_column=c0, end_row=satir, end_column=c0+genis-1)
        st(ws.cell(satir, c0, bas), True, 9, WHITE, NAVY2, "center")
        ws.merge_cells(start_row=satir+1, start_column=c0, end_row=satir+3, end_column=c0+genis-1)
        st(ws.cell(satir+1, c0, ""), False, 9, INK, WHITE, "center")
        ws.merge_cells(start_row=satir+4, start_column=c0, end_row=satir+4, end_column=c0+genis-1)
        st(ws.cell(satir+4, c0, kim), False, 8, GREY, LIGHT, "center")
    ws.row_dimensions[satir+1].height = 22
    ws.row_dimensions[satir+2].height = 22
    return satir+6

def sayfa_ayari(ws, yatay=True, sigdir=1):
    ws.page_setup.orientation = "landscape" if yatay else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = sigdir
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
