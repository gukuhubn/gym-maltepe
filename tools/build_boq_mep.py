# -*- coding: utf-8 -*-
"""Disiplin BoQ'ları — Mekanik ve Elektrik, canlı formüllü .xlsx"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
import proj as P
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

NAVY="FF16273D"; NAVY2="FF24405F"; COPPER="FFB87333"; LIGHT="FFF4F5F7"
CREAM="FFFDF3EC"; WHITE="FFFFFFFF"; RED="FFC8322B"; GIRIS="FFFFF6E8"
thin=Side(style="thin", color="FFD5D8DC"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
TL='#,##0 "TL"'

def st(c,b=False,sz=10,col="FF1C1C1C",fill=None,al="left",wrap=False,fmt=None):
    c.font=Font(name="Calibri",bold=b,size=sz,color=col)
    if fill: c.fill=PatternFill("solid",fgColor=fill)
    c.alignment=Alignment(horizontal=al,vertical="center",wrap_text=wrap)
    c.border=BOX
    if fmt: c.number_format=fmt

def baslik(ws, rng, t, alt=None, sz=14):
    ws.merge_cells(rng); a=rng.split(":")[0]
    st(ws[a],True,sz,WHITE,NAVY,"left"); ws[a]=t
    if alt:
        r2=rng.replace("1","2"); ws.merge_cells(r2)
        b=r2.split(":")[0]; st(ws[b],False,9,"FF6B7078",LIGHT,"left"); ws[b]=alt

def detay_sayfasi(wb, grup, ad):
    ws=wb.create_sheet("2 · Detay metraj")
    baslik(ws,"A1:I1", f"{ad.upper()} — DETAY METRAJ VE KEŞİF",
           f"{P.PROJE} · {P.REV} · {P.TARIH} · Fiyat referansı: {P.FIYAT_TARIH} · "
           f"Birim fiyat sütununu (F) doldurun, tutarlar kendiliğinden hesaplanır")
    for j,x in enumerate(["Poz No","Sistem","Tanım","Birim","Miktar","Birim fiyat (TL)",
                          "Tutar (TL)","Referans düşük (TL)","Referans yüksek (TL)"],1):
        st(ws.cell(3,j,x),True,10,WHITE,NAVY2,"center",True)
    poz=[r for r in P.B if r[1]==grup]
    SIS = {"06":"Havalandırma","06.2":"İklimlendirme","06.3":"Sıhhi tesisat",
           "05.0":"Ana dağıtım","05.1":"Aydınlatma","05.2":"Priz ve kuvvet",
           "05.3":"Zayıf akım","05.4":"Topraklama"}
    def sistem(p):
        if p.startswith("06."):
            n=p[3:5]
            return "Sıhhi tesisat" if n>="30" else ("İklimlendirme" if n>="20" else "Havalandırma")
        n=p[3:5]
        return ("Topraklama" if n>="45" else "Zayıf akım" if n>="30" else
                "Priz ve kuvvet" if n>="20" else "Aydınlatma" if n>="10" else "Ana dağıtım")
    r=4; sistemler={}
    for p_,g,t,br,mik,lo,hi,sen in poz:
        s_=sistem(p_); sistemler.setdefault(s_,[]).append(r)
        alt = CREAM if sen=="—" else (LIGHT if r%2==0 else WHITE)
        st(ws.cell(r,1,p_),True,9,NAVY,alt,"center")
        st(ws.cell(r,2,s_),False,9,"FF6B7078",alt,"left")
        st(ws.cell(r,3,t),False,9,"FF1C1C1C",alt,"left",True)
        st(ws.cell(r,4,br),False,9,"FF1C1C1C",alt,"center")
        st(ws.cell(r,5,round(mik,2)),False,9,"FF1C1C1C",alt,"right",fmt="#,##0.00")
        st(ws.cell(r,6,None),False,10,NAVY,GIRIS,"right",fmt=TL)
        dahil = 0 if sen=="—" else 1
        st(ws.cell(r,7,f"=E{r}*F{r}*{dahil}"),True,9,NAVY,alt,"right",fmt=TL)
        st(ws.cell(r,8,round(mik*lo)*dahil),False,9,"FF8A8F98",alt,"right",fmt=TL)
        st(ws.cell(r,9,round(mik*hi)*dahil),False,9,"FF8A8F98",alt,"right",fmt=TL)
        r+=1
    son=r-1; r+=1
    st(ws.cell(r,3,f"{ad.upper()} İMALAT TOPLAMI"),True,11,WHITE,NAVY,"right")
    for j in (1,2,4,5,6): st(ws.cell(r,j,None),False,10,WHITE,NAVY)
    for j,L_ in ((7,"G"),(8,"H"),(9,"I")):
        st(ws.cell(r,j,f"=SUM({L_}4:{L_}{son})"),True,11,WHITE,NAVY,"right",fmt=TL)
    for col,w in zip("ABCDEFGHI",[9,16,56,8,10,16,16,16,16]): ws.column_dimensions[col].width=w
    ws.freeze_panes="A4"; ws.auto_filter.ref=f"A3:I{son}"
    return sistemler, r

def ozet_sayfasi(wb, grup, ad, sistemler, toplam_satir):
    s1=wb.create_sheet("1 · Özet",0)
    baslik(s1,"A1:E1", f"{ad.upper()} — MALİYET ÖZETİ",
           f"{P.REV} · {P.TARIH} · {P.FIYAT_TARIH} · Ekipman işverence temin edilmiştir, bütçe dışıdır.")
    for j,x in enumerate(["Sistem","Teklif tutarı (TL)","Referans düşük (TL)",
                          "Referans yüksek (TL)","Poz adedi"],1):
        st(s1.cell(4,j,x),True,10,WHITE,NAVY2,"center",True)
    row=5; ilk=row
    for s_,rs in sistemler.items():
        rng=lambda col: "+".join(f"'2 · Detay metraj'!{col}{i}" for i in rs)
        alt = LIGHT if row%2 else WHITE
        st(s1.cell(row,1,s_),True,10,NAVY,alt,"left")
        st(s1.cell(row,2,f"={rng('G')}"),False,10,"FF1C1C1C",alt,"right",fmt=TL)
        st(s1.cell(row,3,f"={rng('H')}"),False,10,"FF8A8F98",alt,"right",fmt=TL)
        st(s1.cell(row,4,f"={rng('I')}"),False,10,"FF8A8F98",alt,"right",fmt=TL)
        st(s1.cell(row,5,len(rs)),False,10,"FF6B7078",alt,"center")
        row+=1
    st(s1.cell(row,1,f"{ad.upper()} İMALAT TOPLAMI"),True,12,WHITE,COPPER,"left")
    for j,L_ in ((2,"B"),(3,"C"),(4,"D")):
        st(s1.cell(row,j,f"=SUM({L_}{ilk}:{L_}{row-1})"),True,12,WHITE,COPPER,"right",fmt=TL)
    st(s1.cell(row,5,sum(len(v) for v in sistemler.values())),True,12,WHITE,COPPER,"center")
    TOP=row; row+=2
    st(s1.cell(row,1,"m² başına (net iç alan 103,78 m²)"),True,10,NAVY,LIGHT,"left")
    for j,L_ in ((2,"B"),(3,"C"),(4,"D")):
        st(s1.cell(row,j,f"={L_}{TOP}/{P.A['ic_toplam']}"),False,10,"FF1C1C1C",LIGHT,"right",fmt=TL)
    st(s1.cell(row,5,None),False,10,"FF1C1C1C",LIGHT); row+=2
    st(s1.cell(row,1,"ANA DOSYA İLE İLİŞKİ"),True,11,WHITE,NAVY2,"left")
    for j in range(2,6): st(s1.cell(row,j,None),True,11,WHITE,NAVY2)
    row+=1
    for t,v in [(f"Bu dosyadaki pozlar, ana BoQ'daki '{grup}' grubunun birebir aynısıdır.",""),
                ("Ana dosya: output/Gym_Maliyet_BoQ.xlsx — 2 · Detay metraj sayfası","")]:
        st(s1.cell(row,1,t),False,9,"FF1C1C1C",WHITE,"left")
        for j in range(2,6): st(s1.cell(row,j,None),False,9,"FF1C1C1C",WHITE)
        row+=1
    for col,w in zip("ABCDE",[44,20,20,20,11]): s1.column_dimensions[col].width=w

def veri_sayfasi(wb, grup):
    ws=wb.create_sheet("3 · Sistem verileri")
    if grup=="MEKANİK":
        baslik(ws,"A1:D1","MEKANİK SİSTEM VERİLERİ","Hava dengesi · kanal · iklimlendirme · boru metrajı")
        blok=[("HAVA DENGESİ",[("Taze hava debisi",P.TAZE,"m³/h"),
               ("Kişi başı taze hava",round(P.TAZE/P.KISI),"m³/h·kişi"),
               ("Hava değişim sayısı",P.ACH,"h⁻¹"),
               ("Salon egzozu",P.TAZE-P.EGZOZ_ISLAK,"m³/h"),
               ("Islak hacim egzozu",P.EGZOZ_ISLAK,"m³/h")]),
              ("KANAL METRAJI",[("Besleme ana hattı 500×150",P.L_KANAL_B,"m"),
               ("Egzoz ana hattı 400×150",P.L_KANAL_E,"m"),
               ("Islak hacim hattı Ø160",P.L_KANAL_I,"m"),
               ("Menfez branşmanı Ø200",P.L_BRANS,"m"),
               ("Valf branşmanı Ø125",P.L_BRANS_I,"m")]),
              ("İKLİMLENDİRME",[("Hesaplanan soğutma yükü",P.SOGUTMA_BTU,"BTU/h"),
               ("Kurulu kapasite",P.KLIMA_BTU,"BTU/h"),
               ("Kapasite marjı",P.SOGUTMA_MARJ,"%"),
               ("İç ünite adedi",P.ADET_KLIMA,"adet"),
               ("Bakır hat uzunluğu",P.L_BAKIR,"m"),
               ("Drenaj hattı uzunluğu",P.L_DRENAJ,"m")]),
              ("SIHHİ TESİSAT",[("Temiz su Ø25",P.L_TEMIZ25,"m"),("Temiz su Ø20",P.L_TEMIZ20,"m"),
               ("Sıcak su Ø20 izoleli",P.L_SICAK,"m"),("Pis su Ø100",P.L_PIS100,"m"),
               ("Pis su Ø70",P.L_PIS70,"m"),("Pis su Ø50",P.L_PIS50,"m")])]
    else:
        baslik(ws,"A1:D1","ELEKTRİK SİSTEM VERİLERİ","Linye tablosu · faz dengesi · yük özeti")
        blok=[("GÜÇ ÖZETİ",[("Toplam bağlı güç",P.BAGLI_KW,"kW"),
               ("Talep gücü (eşzamanlılık)",P.TALEP_KW,"kW"),
               ("Faz dengesizliği",P.FAZ_DENGE,"%"),
               ("Ana kesici",P.ANA_KESICI//3,"×3 A"),
               ("Linye adedi",len(P.LINYE),"adet")]),
              ("FAZ YÜKLERİ",[(f"Faz {f}",P.FAZ_YUK[f],"kW") for f in ("L1","L2","L3")]+
                             [(f"Akım {f}",P.AKIM_FAZ[f],"A") for f in ("L1","L2","L3")]),
              ("CİHAZ ADETLERİ",[("Lineer LED armatür",P.ARMATUR_ADET,"adet"),
               ("IP44 downlight",P.DOWNLIGHT_ADET,"adet"),("İkili topraklı priz",len(P.PRIZ),"adet"),
               ("IP44 priz",len(P.PRIZ_IP44),"adet"),("IP kamera",len(P.KAMERA),"adet"),
               ("Tavan hoparlörü",len(P.HOPARLOR),"adet"),("Duman dedektörü",len(P.DEDEKTOR),"adet")]),
              ("KABLO METRAJI",[("Aydınlatma linyesi 3×1,5",P.L_LINYE15,"m"),
               ("Priz/kuvvet linyesi 3×2,5",P.L_LINYE25,"m"),
               ("Zayıf akım kablolaması",P.L_ZAYIF,"m"),("Kolon hattı NYY 5×10",22,"m")])]
    r=4
    for bas,satirlar in blok:
        st(ws.cell(r,1,bas),True,11,WHITE,NAVY2,"left")
        for j in range(2,5): st(ws.cell(r,j,None),True,11,WHITE,NAVY2)
        r+=1
        for t,v,b in satirlar:
            alt=LIGHT if r%2 else WHITE
            st(ws.cell(r,1,t),False,10,"FF1C1C1C",alt,"left")
            st(ws.cell(r,2,v),False,10,NAVY,alt,"right",fmt="#,##0.00")
            st(ws.cell(r,3,b),False,10,"FF6B7078",alt,"center")
            st(ws.cell(r,4,None),False,10,"FF1C1C1C",alt)
            r+=1
        r+=1
    if grup=="ELEKTRİK":
        st(ws.cell(r,1,"LİNYE TABLOSU"),True,11,WHITE,NAVY2,"left")
        for j in range(2,5): st(ws.cell(r,j,None),True,11,WHITE,NAVY2)
        r+=1
        for j,x in enumerate(["Linye","Tanım","Koruma / kesit","Faz · bağlı kW · talep kW"],1):
            st(ws.cell(r,j,x),True,10,WHITE,NAVY,"center",True)
        r+=1
        for l in P.LINYE:
            alt=LIGHT if r%2 else WHITE
            st(ws.cell(r,1,l[0]),True,10,NAVY,alt,"center")
            st(ws.cell(r,2,l[1]),False,10,"FF1C1C1C",alt,"left",True)
            st(ws.cell(r,3,f"{l[2]} · {l[3]} mm²"),False,10,"FF1C1C1C",alt,"center")
            st(ws.cell(r,4,f"{l[4]} · {l[5]:.2f} · {l[7]:.2f}"),False,10,"FF1C1C1C",alt,"center")
            r+=1
    for col,w in zip("ABCD",[42,16,12,34]): ws.column_dimensions[col].width=w

def varsayim_sayfasi(wb, grup):
    ws=wb.create_sheet("4 · Varsayımlar")
    baslik(ws,"A1:C1","VARSAYIMLAR VE DOĞRULANACAKLAR","Bu değerler ölçülmemiştir")
    ortak=[("Tavan yüksekliği","3,20 m","VARSAYIM — iç hacim, kanal güzergâhı ve armatür adedinin girdisi"),
           ("Geometri doğruluğu","±%3","Raster paftadan vektörleştirme — DXF (R2010) ile kesinleştirilmeli"),
           ("Net iç kullanım alanı","103,78 m²","Pafta m² etiketlerinden kalibre")]
    ozel = ([("Mevcut pis su bağlantı kotu","BİLİNMİYOR","Seçenek A / B kararını belirler — söküm sonrası ilk iş"),
             ("Doğalgaz bağlantısı","Yok varsayıldı","Varsa sıcak su çözümü kombiye revize edilir"),
             ("Havalandırma bacası","9,5 m varsayım","Çatı kotu ölçülecek; bina bacası varsa kullanılabilir"),
             ("Dış ünite montaj yüzeyi","Arka cephe duvarı","Taşıyıcılık ve yönetim onayı doğrulanacak"),
             ("Kanal güzergâhı","Şematik","Kolon ve kiriş konumları yerinde tespit edilecek")]
            if grup=="MEKANİK" else
            [("Mevcut pano gücü / trifaze","BİLİNMİYOR","Talep 17,03 kW — 3×32 A gerekir; yetersizse 3–8 hafta güç artırımı"),
             ("Temel topraklaması","Yok varsayıldı","Varsa 05.45 pozu iptal edilir"),
             ("Kablo metrajları","Güzergâh uzunluğundan","Yaklaşık değer — DXF sonrası kesinleşir"),
             ("Sıcak su çözümü","2 × 6 kW ani ısıtıcı","Boylere çevrilirse talep gücü ≈5 kW düşer, 3×25 A yeterli olur"),
             ("Kamera kapsamı","Soyunma ve WC HARİÇ","Kişisel verilerin korunması — kapsam dışı bırakılmıştır")])
    for j,x in enumerate(["Parametre","Değer","Not"],1):
        st(ws.cell(3,j,x),True,10,WHITE,NAVY2,"center",True)
    r=4
    for t,v,n in ortak+ozel:
        alt=CREAM if v in ("BİLİNMİYOR",) else (LIGHT if r%2 else WHITE)
        st(ws.cell(r,1,t),True,10,NAVY,alt,"left")
        st(ws.cell(r,2,v),False,10,"FF1C1C1C",alt,"right")
        st(ws.cell(r,3,n),False,10,"FF1C1C1C",alt,"left",True); r+=1
    for col,w in zip("ABC",[34,22,78]): ws.column_dimensions[col].width=w

def uret(grup, ad, dosya):
    wb=Workbook()
    sistemler,tr = detay_sayfasi(wb, grup, ad)
    ozet_sayfasi(wb, grup, ad, sistemler, tr)
    veri_sayfasi(wb, grup); varsayim_sayfasi(wb, grup)
    del wb["Sheet"]; wb.save(dosya)
    n=sum(len(v) for v in sistemler.values())
    print(f"→ {dosya}  ·  {n} poz  ·  {len(wb.sheetnames)} sayfa")

if __name__=="__main__":
    uret("MEKANİK","Mekanik tesisat","output/Gym_Mekanik_BoQ.xlsx")
    uret("ELEKTRİK","Elektrik","output/Gym_Elektrik_BoQ.xlsx")
