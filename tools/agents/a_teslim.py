# -*- coding: utf-8 -*-
"""TESLİM DENETİM AJANI — çıktı dosyaları eksiksiz ve okunabilir mi."""
import zipfile
from pathlib import Path
from .base import Ajan
import proj as P

ROOT = Path(__file__).resolve().parent.parent.parent
OUT  = ROOT/"output"

# (dosya, asgari KB, asgari birim [pdf sayfa / xlsx sayfa / zip dosya], açıklama)
# Boyut yalnız "tamamen boş mu" testidir; içerik ölçüsü PAFTA/SAYFA SAYISIDIR.
BEKLENEN = [
 ("Gym_Donusum_Dosyasi_A3.pdf",    60, 12, "Ana dosya — A3 yatay"),
 ("Gym_Mimari_Proje_A3.pdf",       60, 14, "Mimari uygulama seti"),
 ("Gym_Mekanik_Proje_A3.pdf",      60,  9, "Mekanik uygulama seti"),
 ("Gym_Elektrik_Proje_A3.pdf",     60,  8, "Elektrik uygulama seti"),
 ("Gym_ADP_Tek_Hat_Semasi.pdf",    40,  8, "ADP tek hat şeması"),
 ("Gym_Insaat_Seti_A3.pdf",       150, 33, "İnşaat seti — plan/kesit/görünüş/detay"),
 ("Gym_CAD_Paftalar.pdf",         150, 15, "CAD pafta önizleme (A1)"),
 ("Gym_Sunum_16x9.pdf",            60, 12, "Yönetim sunumu"),
 ("Gym_Denetim_Raporu.pdf",        20,  7, "Ajan denetim raporu"),
 ("Gym_Maliyet_BoQ.xlsx",          15,  4, "Ana BoQ — canlı formüllü"),
 ("Gym_Kesif_Ozeti_BoQ.xlsx",      15,  7, "Keşif özeti"),
 ("Gym_Hakedis_Sablonu.xlsx",      10,  5, "Hakediş şablonu"),
 ("Gym_Butce_Takip.xlsx",           8,  2, "Bütçe takip ve nakit akışı"),
 ("Gym_Mekanik_BoQ.xlsx",           8,  4, "Mekanik BoQ"),
 ("Gym_Elektrik_BoQ.xlsx",          8,  4, "Elektrik BoQ"),
 ("Gym_Pano_Yukleme_Cetveli.xlsx",  8,  2, "ADP yükleme cetveli"),
 ("Gym_CAD_Seti_DXF.zip",         200, 19, "DXF seti"),
 ("Gym_Proje_Paketi.zip",        4000, 39, "TEK TESLİM DOSYASI — tüm set"),
 ("Gym_Model.html",               100,  0, "3B model"),
]


def _birim(f):
    """Dosya türüne göre içerik birimi sayısı."""
    u = f.suffix.lower()
    try:
        if u == ".pdf":
            import pypdfium2 as pdfium
            return len(pdfium.PdfDocument(str(f))), "pafta"
        if u == ".xlsx":
            import openpyxl
            return len(openpyxl.load_workbook(f, read_only=True).sheetnames), "sayfa"
        if u == ".zip":
            with zipfile.ZipFile(f) as z:
                return len([n for n in z.namelist() if not n.endswith("/")]), "dosya"
    except Exception:
        return None, ""
    return None, ""


class TeslimAjani(Ajan):
    ad = "teslim"
    baslik = "Teslim paketi — dosya bütünlüğü ve sürüm denetimi"

    def denetle(self, r):
        eksik = 0
        for ad, kb, birim_min, aciklama in BEKLENEN:
            f = OUT/ad
            if not f.exists():
                r.hata("dosya", f"{ad} üretilmemiş — {aciklama}", konum=ad)
                eksik += 1; continue
            boyut = f.stat().st_size/1024
            n, bad = _birim(f)
            if boyut < kb:
                r.hata("dosya", f"{ad} yalnız {boyut:.0f} KB — dosya boş görünüyor",
                       konum=ad)
            elif birim_min and n is not None and n < birim_min:
                r.hata("dosya", f"{ad}: {n} {bad} — beklenen en az {birim_min} {bad}",
                       konum=ad, oneri="İlgili build betiğini yeniden çalıştırın.")
            else:
                ek = f" · {n} {bad}" if n else ""
                r.bilgi("dosya", f"{ad} · {boyut/1024:.2f} MB{ek} — {aciklama}")
        # DXF seti pafta başına ayrı dosya içeriyor mu
        z = OUT/"Gym_CAD_Seti_DXF.zip"
        if z.exists():
            with zipfile.ZipFile(z) as zf:
                dxf = [n for n in zf.namelist() if n.lower().endswith(".dxf")]
            tekil = [n for n in dxf if "/paftalar/" in n or n.count("-") >= 2]
            r.bilgi("dxf", f"DXF setinde {len(dxf)} dosya ({len(tekil)} tekil pafta)")
            if len(dxf) < 5:
                r.uyari("dxf", f"DXF setinde yalnız {len(dxf)} dosya — her pafta ayrı "
                               f"dosya olarak verilmeli")
        # render
        rn = list((OUT/"render").glob("*.png")) if (OUT/"render").exists() else []
        if not rn:
            r.uyari("render", "output/render altında görsel yok")
        else:
            r.bilgi("render", f"{len(rn)} render görseli")
        r.bilgi("sürüm", f"{P.REV} · {P.TARIH} · proje no {P.PROJE.get('no', '—')}"
                if isinstance(P.PROJE, dict) else f"{P.REV} · {P.TARIH}")
        if not eksik:
            r.bilgi("dosya", f"{len(BEKLENEN)} beklenen çıktının tamamı üretildi")
