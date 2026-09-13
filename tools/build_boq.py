# -*- coding: utf-8 -*-
"""TESLİMAT C — BoQ / keşif-metraj, canlı formüllü .xlsx"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
import proj as P
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY="FF16273D"; NAVY2="FF24405F"; COPPER="FFB87333"; LIGHT="FFF4F5F7"
CREAM="FFFDF3EC"; WHITE="FFFFFFFF"; RED="FFC8322B"
thin=Side(style="thin", color="FFD5D8DC")
BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
def st(c, b=False, sz=10, col="FF1C1C1C", fill=None, al="left", wrap=False, fmt=None):
    c.font=Font(name="Calibri", bold=b, size=sz, color=col)
    if fill: c.fill=PatternFill("solid", fgColor=fill)
    c.alignment=Alignment(horizontal=al, vertical="center", wrap_text=wrap)
    c.border=BOX
    if fmt: c.number_format=fmt
TL='#,##0 "TL"'

wb=Workbook()

# ─────────────────────────── 2. DETAY METRAJ (önce kur, özet ona bakacak) ─────
ws=wb.create_sheet("2 · Detay metraj")
ws.merge_cells("A1:K1"); st(ws["A1"],True,14,WHITE,NAVY,"left")
ws["A1"]="DETAY METRAJ VE KEŞİF — birim fiyat sütunlarını doldurun, tutarlar kendiliğinden hesaplanır"
ws.merge_cells("A2:K2"); st(ws["A2"],False,9,"FF6B7078",LIGHT,"left")
ws["A2"]=(f"{P.PROJE} · {P.REV} · {P.TARIH} · Fiyat referansı: {P.FIYAT_TARIH} · "
          f"Net iç alan {P.A['ic_toplam']} m²")
hdr=["Poz No","Grup","Tanım","Birim","Miktar","Birim fiyat (TL)",
     "Tutar (TL)","Referans düşük (TL)","Referans yüksek (TL)","Senaryo","Dâhil"]
for j,x in enumerate(hdr,1):
    st(ws.cell(3,j,x),True,10,WHITE,NAVY2,"center",True)
r=4; grup_satir={}
for poz,grup,tanim,birim,mik,lo,hi,sen in P.B:
    if grup not in grup_satir: grup_satir[grup]=[]
    grup_satir[grup].append(r)
    alt = LIGHT if r%2==0 else WHITE
    if sen=="—" or poz in ("03.08","03.09"): alt=CREAM
    st(ws.cell(r,1,poz),True,9,"FF16273D",alt,"center")
    st(ws.cell(r,2,grup),False,9,"FF6B7078",alt,"left")
    st(ws.cell(r,3,tanim),False,9,"FF1C1C1C",alt,"left",True)
    st(ws.cell(r,4,birim),False,9,"FF1C1C1C",alt,"center")
    st(ws.cell(r,5,round(mik,2)),False,9,"FF1C1C1C",alt,"right",fmt="#,##0.00")
    st(ws.cell(r,6,None),False,10,"FF16273D","FFFFF6E8","right",fmt=TL)   # BOŞ — kullanıcı girer
    st(ws.cell(r,7,f"=E{r}*F{r}*K{r}"),True,9,"FF16273D",alt,"right",fmt=TL)
    st(ws.cell(r,8,f"={round(mik*lo)}*K{r}"),False,9,"FF8A8F98",alt,"right",fmt=TL)
    st(ws.cell(r,9,f"={round(mik*hi)}*K{r}"),False,9,"FF8A8F98",alt,"right",fmt=TL)
    kod = "A" if poz=="03.08" else ("B" if poz=="03.09" else sen)
    st(ws.cell(r,10,kod),False,9,"FF6B7078",alt,"center")
    fml = {"M":"1",
           "O":"=IF('1 · Özet'!$B$4=\"ÖNERİLEN\",1,0)",
           "A":"=IF('1 · Özet'!$E$4=\"A\",1,0)",
           "B":"=IF('1 · Özet'!$E$4=\"B\",1,0)"}[kod]
    st(ws.cell(r,11,1 if fml=="1" else fml),False,9,"FF6B7078",alt,"center")
    r+=1
son=r-1
r+=1
st(ws.cell(r,3,"İMALAT ARA TOPLAMI"),True,11,WHITE,NAVY,"right")
for j in (1,2,4,5,6,10,11): st(ws.cell(r,j,None),False,10,WHITE,NAVY)
st(ws.cell(r,7,f"=SUM(G4:G{son})"),True,11,WHITE,NAVY,"right",fmt=TL)
st(ws.cell(r,8,f"=SUM(H4:H{son})"),True,11,WHITE,NAVY,"right",fmt=TL)
st(ws.cell(r,9,f"=SUM(I4:I{son})"),True,11,WHITE,NAVY,"right",fmt=TL)
ARA=r
for col,w in zip("ABCDEFGHIJK",[9,17,52,8,10,16,16,16,16,10,8]): ws.column_dimensions[col].width=w
ws.freeze_panes="A4"; ws.auto_filter.ref=f"A3:K{son}"

# ─────────────────────────── 1. ÖZET ──────────────────────────────────────────
s1=wb.create_sheet("1 · Özet", 0)
s1.merge_cells("A1:F1"); st(s1["A1"],True,16,WHITE,NAVY,"left")
s1["A1"]="MALTEPE / İDEALTEPE — GYM DÖNÜŞÜMÜ · MALİYET ÖZETİ"
s1.merge_cells("A2:F2"); st(s1["A2"],False,9,"FF6B7078",LIGHT,"left")
s1["A2"]=f"{P.REV} · {P.TARIH} · {P.FIYAT_TARIH} · Ekipman işverence temin edilmiştir, bütçe dışıdır."
st(s1["A4"],True,11,WHITE,NAVY2,"left"); s1["A4"]="SENARYO SEÇİMİ"
st(s1["B4"],True,11,"FF16273D","FFFFF6E8","center"); s1["B4"]="ÖNERİLEN"
dv=DataValidation(type="list", formula1='"MİNİMUM,ÖNERİLEN"', allow_blank=False)
s1.add_data_validation(dv); dv.add(s1["B4"])
st(s1["C4"],False,9,"FF6B7078",WHITE,"left")
s1["C4"]="Açılır listeden seçin — tüm toplamlar anında yeniden hesaplanır"
st(s1["D4"],True,11,WHITE,NAVY2,"left"); s1["D4"]="ISLAK HACİM"
st(s1["E4"],True,11,"FF16273D","FFFFF6E8","center"); s1["E4"]="A"
dv2=DataValidation(type="list", formula1='"A,B"', allow_blank=False)
s1.add_data_validation(dv2); dv2.add(s1["E4"])
st(s1["F4"],False,9,"FF6B7078",WHITE,"left"); s1["F4"]="A: zemin yükseltme · B: pompa"
row=6
for j,x in enumerate(["Maliyet başlığı","Teklif tutarı (TL)","Referans düşük (TL)",
                      "Referans yüksek (TL)","Poz adedi","Pay (%)"],1):
    st(s1.cell(row,j,x),True,10,WHITE,NAVY2,"center",True)
row=7; ilk=row
for g in P.GRUPLAR:
    rs=grup_satir[g]
    rng=lambda col: "+".join(f"'2 · Detay metraj'!{col}{i}" for i in rs)
    st(s1.cell(row,1,g),True,10,"FF16273D",LIGHT if row%2 else WHITE,"left")
    st(s1.cell(row,2,f"={rng('G')}"),False,10,"FF1C1C1C",LIGHT if row%2 else WHITE,"right",fmt=TL)
    st(s1.cell(row,3,f"={rng('H')}"),False,10,"FF8A8F98",LIGHT if row%2 else WHITE,"right",fmt=TL)
    st(s1.cell(row,4,f"={rng('I')}"),False,10,"FF8A8F98",LIGHT if row%2 else WHITE,"right",fmt=TL)
    st(s1.cell(row,5,len(rs)),False,10,"FF6B7078",LIGHT if row%2 else WHITE,"center")
    st(s1.cell(row,6,f"=IFERROR((C{row}+D{row})/2/(($C${ilk+len(P.GRUPLAR)}+$D${ilk+len(P.GRUPLAR)})/2),0)"),
       False,10,"FF6B7078",LIGHT if row%2 else WHITE,"right",fmt="0%")
    row+=1
st(s1.cell(row,1,"İMALAT ARA TOPLAMI"),True,11,WHITE,NAVY,"left")
for col,L_ in ((2,"B"),(3,"C"),(4,"D")):
    st(s1.cell(row,col,f"=SUM({L_}{ilk}:{L_}{row-1})"),True,11,WHITE,NAVY,"right",fmt=TL)
st(s1.cell(row,5,len(P.B)),True,11,WHITE,NAVY,"center")
st(s1.cell(row,6,1),True,11,WHITE,NAVY,"right",fmt="0%")
IM=row; row+=1
for ad,oran in (("Şantiye genel giderleri",P.V["santiye_gider"][0]),
                ("Beklenmedik giderler",P.V["beklenmedik"][0])):
    st(s1.cell(row,1,f"{ad} (%{oran*100:.0f})"),True,10,"FF16273D",CREAM,"left")
    baz = f"B{IM}" if row==IM+1 else f"(B{IM}+B{IM+1})"
    bazC= f"C{IM}" if row==IM+1 else f"(C{IM}+C{IM+1})"
    bazD= f"D{IM}" if row==IM+1 else f"(D{IM}+D{IM+1})"
    st(s1.cell(row,2,f"={baz}*{oran}"),False,10,"FF1C1C1C",CREAM,"right",fmt=TL)
    st(s1.cell(row,3,f"={bazC}*{oran}"),False,10,"FF8A8F98",CREAM,"right",fmt=TL)
    st(s1.cell(row,4,f"={bazD}*{oran}"),False,10,"FF8A8F98",CREAM,"right",fmt=TL)
    st(s1.cell(row,5,None),False,10,"FF1C1C1C",CREAM); st(s1.cell(row,6,None),False,10,"FF1C1C1C",CREAM)
    row+=1
st(s1.cell(row,1,"GENEL TOPLAM"),True,13,WHITE,COPPER,"left")
for col,L_ in ((2,"B"),(3,"C"),(4,"D")):
    st(s1.cell(row,col,f"={L_}{IM}+{L_}{IM+1}+{L_}{IM+2}"),True,13,WHITE,COPPER,"right",fmt=TL)
st(s1.cell(row,5,None),True,13,WHITE,COPPER); st(s1.cell(row,6,None),True,13,WHITE,COPPER)
TOP=row; row+=2
st(s1.cell(row,1,"m² başına (net iç alan)"),True,10,"FF16273D",LIGHT,"left")
for col,L_ in ((2,"B"),(3,"C"),(4,"D")):
    st(s1.cell(row,col,f"={L_}{TOP}/{P.A['ic_toplam']}"),False,10,"FF1C1C1C",LIGHT,"right",fmt=TL)
row+=2
st(s1.cell(row,1,"AÇILIŞ ÖNCESİ NAKİT İHTİYACI"),True,12,WHITE,NAVY2,"left")
for j in range(2,7): st(s1.cell(row,j,None),True,12,WHITE,NAVY2)
row+=1
st(s1.cell(row,1,"Tadilat (genel toplam)"),False,10,"FF1C1C1C",WHITE,"left")
st(s1.cell(row,2,f"=B{TOP}"),False,10,"FF1C1C1C",WHITE,"right",fmt=TL)
st(s1.cell(row,3,f"=C{TOP}"),False,10,"FF8A8F98",WHITE,"right",fmt=TL)
st(s1.cell(row,4,f"=D{TOP}"),False,10,"FF8A8F98",WHITE,"right",fmt=TL)
nb=row; row+=1
for ad,lo,hi in P.NAKIT[1:]:
    st(s1.cell(row,1,ad),False,10,"FF1C1C1C",LIGHT if row%2 else WHITE,"left")
    st(s1.cell(row,2,None),False,10,"FF16273D","FFFFF6E8","right",fmt=TL)
    st(s1.cell(row,3,lo),False,10,"FF8A8F98",LIGHT if row%2 else WHITE,"right",fmt=TL)
    st(s1.cell(row,4,hi),False,10,"FF8A8F98",LIGHT if row%2 else WHITE,"right",fmt=TL)
    st(s1.cell(row,5,None),False,10,"FF1C1C1C",LIGHT if row%2 else WHITE)
    st(s1.cell(row,6,None),False,10,"FF1C1C1C",LIGHT if row%2 else WHITE)
    row+=1
st(s1.cell(row,1,"TOPLAM AÇILIŞ ÖNCESİ NAKİT"),True,12,WHITE,COPPER,"left")
for col,L_ in ((2,"B"),(3,"C"),(4,"D")):
    st(s1.cell(row,col,f"=SUM({L_}{nb}:{L_}{row-1})"),True,12,WHITE,COPPER,"right",fmt=TL)
st(s1.cell(row,5,None),True,12,WHITE,COPPER); st(s1.cell(row,6,None),True,12,WHITE,COPPER)
for col,w in zip("ABCDEF",[46,20,20,20,11,10]): s1.column_dimensions[col].width=w

# ─────────────────────────── 3. EKİPMAN ───────────────────────────────────────
s3=wb.create_sheet("3 · Ekipman (işverence temin)")
s3.merge_cells("A1:H1"); st(s3["A1"],True,14,WHITE,NAVY,"left")
s3["A1"]="EKİPMAN BİLGİ SAYFASI — BÜTÇE DIŞI"
s3.merge_cells("A2:H2"); st(s3["A2"],False,9,"FF6B7078",CREAM,"left")
s3["A2"]=("Ekipman işverence satın alınmıştır; bu dosyadaki maliyet kalemlerine dâhil değildir. "
          "Marka/model listesi geldiğinde aşağıdaki sütunlar doldurulmalıdır.")
for j,x in enumerate(["Kod","Ekipman","En (cm)","Boy (cm)","Yükseklik (m)","Adet",
                      "Footprint (m²)","Marka / model (işveren dolduracak)"],1):
    st(s3.cell(3,j,x),True,10,WHITE,NAVY2,"center",True)
r=4
for kod,ad,en,boy,adet,_,_,tip,hh in P.EKIPMAN:
    alt=LIGHT if r%2==0 else WHITE
    st(s3.cell(r,1,kod),True,10,"FF16273D",alt,"center")
    st(s3.cell(r,2,ad),False,10,"FF1C1C1C",alt,"left",True)
    st(s3.cell(r,3,en if en else "altıgen"),False,10,"FF1C1C1C",alt,"right")
    st(s3.cell(r,4,boy if boy else "10,60 m²"),False,10,"FF1C1C1C",alt,"right")
    st(s3.cell(r,5,hh),False,10,"FF1C1C1C",alt,"right",fmt="#,##0.00")
    st(s3.cell(r,6,adet),False,10,"FF1C1C1C",alt,"center")
    fp = P.HEX_M2 if en is None else round(en*boy/10000*adet,2)
    st(s3.cell(r,7,fp),False,10,"FF1C1C1C",alt,"right",fmt="#,##0.00")
    st(s3.cell(r,8,None),False,10,"FF16273D","FFFFF6E8","left")
    r+=1
st(s3.cell(r,2,"TOPLAM FOOTPRINT"),True,11,WHITE,NAVY,"right")
for j in (1,3,4,5,8): st(s3.cell(r,j,None),True,11,WHITE,NAVY)
st(s3.cell(r,6,f"=SUM(F4:F{r-1})"),True,11,WHITE,NAVY,"center")
st(s3.cell(r,7,f"=SUM(G4:G{r-1})"),True,11,WHITE,NAVY,"right",fmt="#,##0.00")
r+=2
st(s3.cell(r,2,"Salon brüt alanı"),False,10,"FF1C1C1C",LIGHT,"left")
st(s3.cell(r,7,P.A["salon"]),False,10,"FF1C1C1C",LIGHT,"right",fmt="#,##0.00"); r+=1
st(s3.cell(r,2,"Serbest sirkülasyon / çalışma alanı"),False,10,"FF1C1C1C",LIGHT,"left")
st(s3.cell(r,7,f"={P.A['salon']}-G{r-3}"),False,10,"FF1C1C1C",LIGHT,"right",fmt="#,##0.00")
for col,w in zip("ABCDEFGH",[8,44,12,12,14,8,16,40]): s3.column_dimensions[col].width=w

# ─────────────────────────── 4. VARSAYIMLAR ───────────────────────────────────
s4=wb.create_sheet("4 · Varsayımlar")
s4.merge_cells("A1:D1"); st(s4["A1"],True,14,WHITE,NAVY,"left")
s4["A1"]="VARSAYIMLAR, METRAJ DAYANAĞI VE DOĞRULANACAKLAR"
for j,x in enumerate(["Parametre","Değer","Birim","Kaynak / varsayım notu"],1):
    st(s4.cell(3,j,x),True,10,WHITE,NAVY2,"center",True)
r=4
veri=[("Salon alanı",P.A["salon"],"m²","Alan Dağılımı paftası etiketi"),
      ("Erkek soyunma bloğu",P.A["erkek_blok"],"m²","Alan Dağılımı paftası etiketi"),
      ("Kadın soyunma bloğu",P.A["kadin_blok"],"m²","Alan Dağılımı paftası etiketi"),
      ("Net iç kullanım alanı",P.A["ic_toplam"],"m²","Toplam — tüm metrajın temeli"),
      ("Ön bahçe",P.A["on_bahce"],"m²","Açık alan"),
      ("Arka bahçe",P.A["arka_bahce"],"m²","Açık alan"),
      ("Salon dış çevre uzunluğu",P.L_SALON,"m","Vektörleştirilmiş poligon"),
      ("Islak hacim blok çevresi",P.L_ISLAK,"m","Vektörleştirilmiş poligon"),
      ("Boyanacak / sıvalanacak duvar",P.M2_DUVAR_SALON,"m²","Çevre × tavan − cephe doğraması"),
      ("Islak hacim duvar fayansı",P.M2_SERAMIK_D,"m²","Blok çevresi × 2,20 m"),
      ("Yeni alçıpan bölme",P.M2_YENI_BOLME,"m²","Blok içi bölme payı × tavan"),
      ("Armatür adedi",P.ARMATUR_ADET,"adet","Lux hesabı — bakım 0,80 · verim 0,70 · 4400 lm"),
      ("Klima adedi",P.ADET_KLIMA,"adet",f"{P.SOGUTMA_BTU} BTU / 24.000"),
      ("Taze hava debisi",P.TAZE,"m³/h","max(kişi başı 55 m³/h; 3 hava değişimi/saat)"),
      ("Kurulu güç",P.KURULU_KW,"kW","Yük tablosu toplamı"),
      ("Şantiye genel gideri",P.V["santiye_gider"][0],"oran","%8–10 ortası"),
      ("Beklenmedik pay",P.V["beklenmedik"][0],"oran","Islak hacim belirsizliği yüksek")]
for ad,vv,br,kn in veri:
    alt=LIGHT if r%2==0 else WHITE
    st(s4.cell(r,1,ad),True,10,"FF16273D",alt,"left")
    st(s4.cell(r,2,vv),False,10,"FF1C1C1C",alt,"right",fmt="#,##0.00")
    st(s4.cell(r,3,br),False,10,"FF6B7078",alt,"center")
    st(s4.cell(r,4,kn),False,10,"FF1C1C1C",alt,"left",True); r+=1
r+=1
st(s4.cell(r,1,"DOĞRULANACAK — bu değerler ölçülmemiştir"),True,12,WHITE,RED,"left")
for j in (2,3,4): st(s4.cell(r,j,None),True,12,WHITE,RED)
r+=1
for ad,vv,br,kn in [("Tavan yüksekliği",P.V["tavan_h"][0],"m",P.V["tavan_h"][2]),
                    ("Mevcut pis su kotu","—","cm","BİLİNMİYOR — Seçenek A/B kararını belirler"),
                    ("Elektrik pano gücü","—","kW","BİLİNMİYOR — talep 12,1 kW"),
                    ("Üst katta konut","VAR","E/H","VARSAYIM (konservatif) — akustik kalemleri buna bağlı"),
                    ("Doğalgaz bağlantısı","—","E/H","Varsa sıcak su çözümü kombiye döner (poz 03.07)"),
                    ("Geometri doğruluğu","±3","%","Raster pafta izleme — DXF R2010 istenmeli")]:
    st(s4.cell(r,1,ad),True,10,"FF16273D",CREAM,"left")
    st(s4.cell(r,2,vv),False,10,"FF1C1C1C",CREAM,"right")
    st(s4.cell(r,3,br),False,10,"FF6B7078",CREAM,"center")
    st(s4.cell(r,4,kn),False,10,"FF1C1C1C",CREAM,"left",True); r+=1
for col,w in zip("ABCD",[34,14,10,66]): s4.column_dimensions[col].width=w

del wb["Sheet"]
wb.save("output/Gym_Maliyet_BoQ.xlsx")
print("→ output/Gym_Maliyet_BoQ.xlsx  ·", len(P.B), "poz ·", len(wb.sheetnames), "sayfa")
